"""路线对比（D）、导出报告（F）、回填实验结果（G）、校准（M14）。

这个文件测的是什么
    一次分析跑完以后，用户接着会做的三件事。接口都在 syntherisk/api/reports.py：
        D1 换权重对比、D5 差异步骤   POST  /api/analyses/{id}/compare                  compare
        D4 导出对比表                POST  /api/analyses/{id}/compare/export           export_comparison
        F1–F3 生成报告               POST  /api/analyses/{id}/reports                  create_report
              下载                   GET   /api/reports/{id}/download                  download_report
        物料清单预览                 GET   /api/analyses/{id}/reports/materials        preview_materials
        G1–G4 记录实验结果           POST  /api/analyses/{id}/routes/{route_id}/steps/{n}/experiments
                                                                                       record_experiment
        这次分析的回填记录           GET   /api/analyses/{id}/experiments              list_experiments
        G3 撤回或补上授权            PATCH /api/experiments/{id}                       update_experiment
        G2 附件                      POST  /api/experiments/{id}/attachments           add_attachment
              下载                   GET   /api/experiments/attachments/{id}           download_attachment
    另外用到项目历史 GET /api/projects/{id}/history（syntherisk/api/core.py 的 project_history）。
    背后的模块：
        M10 评分       syntherisk/modules/m10_scoring.py    对比表的数据（compare_routes）、换权重重算（score_result）
        M13 报告       syntherisk/modules/m13_reports.py    PDF、Word、结构文件、对比表、物料清单
        M14 回填校准   syntherisk/modules/m14_backfill.py   保存记录、授权数据的匿名导出、校准统计
        文件存储       syntherisk/storage.py                报告和附件加密后存盘
    这三件事都不重新分析，用的是已经存好的评估结果。

怎么检查生成的文件
    接口返回的是文件的字节。测试里不落盘，用 io.BytesIO 把字节包成内存里的“文件”，再用对应的库读回来：
    openpyxl 读 Excel，python-docx（import 名是 docx）读 Word，标准库的 zipfile 读压缩包。

共用的夹具和辅助函数（tests/conftest.py）
    client、account、ctx   测试客户端、已注册的账号、运行时上下文
    register()             再注册一个别的课题组的账号
    start_analysis()       提交一次分析；测试里队列是同步的，返回时已经跑完。
                           不指定 smiles 时目标是布洛芬，三条演示路线是 r1 格氏羧基化（4 步）、
                           r2 Boots（6 步）、r3 BHC（3 步），来自 data/demo_routes/ibuprofen.json
    本文件没有替身：报告、表格、压缩包都是后端真的生成出来的。

本文件里有什么
    test_compare_marks_shared_steps_and_reweights_on_request
        对比标出几条路线共有的步骤；给了权重就按新权重重算，不改保存的结果；权重或路线不对被拒绝
    test_comparison_export_downloads_a_spreadsheet
        对比表可以导出成 Excel 或 CSV，列的先后和请求里的路线顺序一致
    test_a_frontend_on_another_origin_can_read_the_download_file_name
        前端和后端不在同一个地址（跨域）时，前端的脚本也读得到下载文件的文件名
    test_reports_in_three_formats_are_saved_to_project_history
        PDF、Word、结构文件三种报告都能生成和下载，只包含选中的部分，免责声明固定有；都记进项目历史
    test_bill_of_materials_scales_with_target_amount
        物料清单按目标产量折算起始原料的用量
    test_recording_an_experiment_marks_the_step_as_done
        回填以后这一步显示“已做”；只有授权了的记录进入校准数据，而且是匿名的；授权可以撤回
    test_experiment_attachments_are_stored_encrypted
        实验附件加密存盘、能原样下载；只收图片和 PDF；别的课题组下载不了
"""

import io
import zipfile

from docx import Document
from openpyxl import load_workbook

from syntherisk.modules import m14_backfill
from tests.conftest import register, start_analysis


def test_compare_marks_shared_steps_and_reweights_on_request(client, account):
    """D1、D5 路线对比：标出几条路线共有的步骤；请求里给了权重就按新权重重算，但不改保存的结果。

    保证的行为：POST /api/analyses/{id}/compare（api/reports.py 的 compare）
      - 每条路线的每一步带 shared 标记：这一步在被对比的每条路线里都有时是 true（m10_scoring.py 的
        compare_routes，按步骤编号求交集），前端据此把共同步骤灰显（D5）；
      - radar 是六个维度各步的平均分（雷达图用）；
      - 请求里带 weights 时，在结果的副本上按这组权重重新评分（api/reports.py 的 _weighted_result），
        没写的维度按 0 算；数据库里存的结果不变；
      - 权重里有不认识的维度返回 422；要对比的路线一条都找不到返回 404。
    怎么读：不带约束跑一次布洛芬的分析，对比 r1（格氏羧基化）和 r2（Boots）。先看默认权重下的对比结果，
    再只给安全一项权重重算一次，最后试两种出错的请求。
    """
    analysis = start_analysis(client, account)
    url = f"/api/analyses/{analysis['id']}/compare"
    comparison = client.post(url, headers=account["headers"], json={"route_ids": ["r1", "r2"]}).json()
    # 返回的路线按分数排，所以先做成 {路线名: 这条路线} 再取
    routes = {route["name"]: route for route in comparison["routes"]}
    assert set(routes) == {"格氏羧基化", "Boots"}
    # 两条路线只有第 1 步（傅-克酰基化）相同，其余都是差异步骤
    # 这两条路线的第 1 步都是异丁基苯 + 乙酸酐、三氯化铝催化，结果里是同一个步骤编号；格氏路线 4 步，Boots 6 步
    assert [step["shared"] for step in routes["格氏羧基化"]["steps"]] == [True, False, False, False]
    assert [step["shared"] for step in routes["Boots"]["steps"]] == [True] + [False] * 5
    # 雷达图的六个轴就是六个评分维度
    assert set(routes["Boots"]["radar"]) == {"literature", "operation", "safety", "availability", "yield", "green"}

    # 只看安全一项时，分数就等于安全维度按路线公式算出来的值
    # 下面实际核对的是：返回值里的权重确实换了（安全是 1，没写的操作是 0），
    # 并且排第一的路线的分数和默认权重时不一样
    safety_only = client.post(url, headers=account["headers"], json={"route_ids": ["r1", "r2"], "weights": {"safety": 1}}).json()
    assert safety_only["weights"]["safety"] == 1 and safety_only["weights"]["operation"] == 0
    assert safety_only["routes"][0]["score"] != comparison["routes"][0]["score"]
    # 换权重只影响这一次对比，保存的结果不变
    # 重新取结果页：安全的权重还是 data/scoring.toml 里的 0.2
    stored = client.get(f"/api/analyses/{analysis['id']}/result", headers=account["headers"]).json()
    assert stored["weights"]["safety"] == 0.2

    # 权重的键不是六个维度之一：422
    assert client.post(url, headers=account["headers"], json={"route_ids": ["r1"], "weights": {"颜值": 1}}).status_code == 422
    # 路线编号不存在：一条能对比的路线都没有，404
    assert client.post(url, headers=account["headers"], json={"route_ids": ["不存在"]}).status_code == 404


def test_comparison_export_downloads_a_spreadsheet(client, account):
    """D4 导出对比表：可以导出成 Excel 或 CSV，内容是一行一个指标、一列一条路线。

    保证的行为：POST /api/analyses/{id}/compare/export（api/reports.py 的 export_comparison）直接返回文件内容，
    文件名放在响应头 Content-Disposition 里；表格由 m13_reports.py 的 comparison_xlsx、comparison_csv 生成：
    第一行是表头（“指标”加各条路线的名字），下面依次是综合分、步数等指标；各条路线的先后和请求里
    route_ids 的顺序一致。CSV 是带 BOM 的 UTF-8，Excel 直接打开不会乱码。
    怎么读：先导出三条路线的 Excel，用 openpyxl 读回来看前几行；再导出两条路线的 CSV，看第一行。
    """
    analysis = start_analysis(client, account)
    response = client.post(f"/api/analyses/{analysis['id']}/compare/export", headers=account["headers"],
                           json={"route_ids": ["r1", "r2", "r3"], "format": "xlsx"})
    # 中文文件名在响应头里写成 filename*=UTF-8''…（百分号编码），这里只看这种写法在不在
    assert response.status_code == 200 and "filename*=UTF-8''" in response.headers["content-disposition"]
    # response.content 是文件的字节；包成内存里的文件交给 openpyxl 打开。active 是第一张表（“路线对比”）
    sheet = load_workbook(io.BytesIO(response.content)).active
    # 把整张表读成“行的列表”，每一行是各个单元格的值
    rows = [[cell.value for cell in row] for row in sheet.iter_rows()]
    # 表头：三条路线按请求里的顺序 r1、r2、r3
    assert rows[0] == ["指标", "格氏羧基化", "Boots", "BHC"]
    # rows[1] 是“综合分”，rows[2] 是“步数”：三条路线分别是 4、6、3 步
    assert rows[2] == ["步数", 4, 6, 3]

    as_csv = client.post(f"/api/analyses/{analysis['id']}/compare/export", headers=account["headers"],
                         json={"route_ids": ["r1", "r3"], "format": "csv"})
    # utf-8-sig 解码时会去掉文件开头的 BOM；splitlines()[0] 是第一行，也就是表头
    assert as_csv.content.decode("utf-8-sig").splitlines()[0] == "指标,格氏羧基化,BHC"


def test_a_frontend_on_another_origin_can_read_the_download_file_name(client, account):
    """跨域时前端的脚本也读得到下载文件的文件名。

    保证的行为：下载文件的接口把文件名放在响应头 Content-Disposition 里。前端页面和后端不在同一个地址时
    （比如前端开发服务器在 http://localhost:5173，后端在 http://127.0.0.1:8000），浏览器默认只让脚本读
    少数几个响应头，Content-Disposition 不在其中；后端要在跨域设置里点名它（syntherisk/main.py 的 create_app
    给 CORSMiddleware 传的 expose_headers），浏览器看到响应头 Access-Control-Expose-Headers 里有它才放行。
    导出对比表的接口只在这个响应头里给文件名，读不到的话前端就不知道该把文件存成什么名字。
    怎么读：请求里带上 Origin 请求头，装作是 http://localhost:5173 上的页面发来的（这个地址在 CORS_ORIGINS 的
    默认值里，见 syntherisk/config.py），然后看返回的三个响应头。
    """
    analysis = start_analysis(client, account)
    # {**a, "Origin": …}：在登录请求头的基础上再加一个 Origin。浏览器跨域发请求时会自动带上它
    headers = {**account["headers"], "Origin": "http://localhost:5173"}
    response = client.post(f"/api/analyses/{analysis['id']}/compare/export", headers=headers,
                           json={"route_ids": ["r1", "r2"], "format": "csv"})
    assert response.status_code == 200
    # 允许这个来源的页面读这次的返回
    assert response.headers["access-control-allow-origin"] == "http://localhost:5173"
    # 并且点名允许脚本读 Content-Disposition
    assert response.headers["access-control-expose-headers"] == "Content-Disposition"
    assert "filename*=UTF-8''" in response.headers["content-disposition"]


def test_reports_in_three_formats_are_saved_to_project_history(client, account):
    """F1–F3 导出报告：PDF、Word、结构文件三种格式都能生成和下载；只包含选中的部分；生成过的都在项目历史里。

    保证的行为：
      - POST /api/analyses/{id}/reports（api/reports.py 的 create_report）返回 201 和一条报告记录，
        里面的 download_url 再 GET 一次才拿到文件（download_report）；
      - 文件名：没指定时是“目标分子_路线名_日期.扩展名”（m13_reports.py 的 default_filename）；
        指定了但没带扩展名时自动补上；
      - F1 选择内容：sections 里没选的部分不出现；报告标题和最后的免责声明不管选了什么都有
        （m13_reports.py 的 build_blocks）；
      - F2 选 sdf 时给的是一个 zip：route.sdf（路线上的全部分子）、每一步一个 step_N.rxn、一个 README.txt
        （m13_reports.py 的 export_structures_zip）；
      - 报告记进 reports 表，GET /api/projects/{id}/history 里能看到，最新的在前；
      - 报告按课题组隔离：别的组拿着下载地址也是 404。
    怎么读：给格氏路线（r1）依次生成三种格式，每种都下载回来看一眼内容；然后看项目历史；
    最后试别的课题组下载、以及给不存在的路线出报告。
    """
    analysis = start_analysis(client, account)
    url = f"/api/analyses/{analysis['id']}/reports"

    # ---- PDF：不指定 sections，就是五个部分全选
    pdf = client.post(url, headers=account["headers"], json={"route_id": "r1", "format": "pdf"})
    assert pdf.status_code == 201, pdf.text
    # 默认文件名的前半截是目标分子名和路线名，中间是当天的日期（这里不比日期）
    assert pdf.json()["filename"].startswith("布洛芬_格氏羧基化_") and pdf.json()["filename"].endswith(".pdf")
    content = client.get(pdf.json()["download_url"], headers=account["headers"]).content
    # PDF 文件都以 %PDF 开头；大小超过一万字节，说明里面确实有内容，不是一个空壳
    assert content.startswith(b"%PDF") and len(content) > 10_000

    # ---- Word：只选“推荐路线”和“物料清单”两部分，并且自己起文件名（没带扩展名）
    word = client.post(url, headers=account["headers"],
                       json={"route_id": "r1", "format": "docx", "sections": ["route", "materials"], "filename": "给导师看的"}).json()
    assert word["filename"] == "给导师看的.docx"
    # 用 python-docx 把下载回来的字节读成文档，把各段的文字用换行连成一整段来查。
    # 标题也算段落；表格里的文字不在 paragraphs 里
    document = Document(io.BytesIO(client.get(word["download_url"], headers=account["headers"]).content))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    # 报告标题，以及选中的两部分的章节标题。章节的序号是固定的，只选一部分时不会重排，所以是“一”和“四”
    assert "布洛芬 合成路线报告" in text and "一、推荐路线" in text and "四、试剂与物料清单" in text
    assert "二、每一步的证据卡" not in text  # 没选的内容不出现
    assert "免责声明" in text and "不是实验操作规程" in text  # 免责声明固定包含

    # ---- 结构文件：format 写 sdf，拿到的是一个 zip
    structures = client.post(url, headers=account["headers"], json={"route_id": "r1", "format": "sdf"}).json()
    archive = zipfile.ZipFile(io.BytesIO(client.get(structures["download_url"], headers=account["headers"]).content))
    # namelist() 是压缩包里的文件名。格氏路线 4 步，所以有 step_1 到 step_4 四个反应文件
    assert sorted(archive.namelist()) == ["README.txt", "route.sdf", "step_1.rxn", "step_2.rxn", "step_3.rxn", "step_4.rxn"]
    sdf = archive.read("route.sdf").decode("utf-8")
    # 目标 1 + 中间体 3 + 起始原料 4 + 试剂 5
    # SDF 文件里每个分子的记录以一行 $$$$ 结束，数它的个数就是分子的个数。每个分子带一个 Role 属性，
    # starting_material 是起始原料。RXN 文件（一步反应的反应物和产物）以 $RXN 开头
    assert sdf.count("$$$$") == 13 and "starting_material" in sdf and "$RXN" in archive.read("step_1.rxn").decode()

    # ---- 项目历史：三份报告都在，最新生成的排在前面；这次分析也在，并且带着排第一的路线
    history = client.get(f"/api/projects/{account['project_id']}/history", headers=account["headers"]).json()
    assert [report["format"] for report in history["reports"]] == ["sdf", "docx", "pdf"]
    assert history["analyses"][0]["best_route"]["name"]

    # 别的课题组拿着同一个下载地址：404
    other = register(client, email="other@example.com", org="隔壁课题组")
    assert client.get(pdf.json()["download_url"], headers=other["headers"]).status_code == 404
    # 这次分析只有 r1、r2、r3 三条路线，没有 r9：404
    assert client.post(url, headers=account["headers"], json={"route_id": "r9"}).status_code == 404


def test_bill_of_materials_scales_with_target_amount(client, account):
    """物料清单：按想得到多少毫克目标物，倒推每个起始原料要多少；试剂只列名字，没有用量。

    保证的行为：GET /api/analyses/{id}/reports/materials（api/reports.py 的 preview_materials，
    计算在 m13_reports.py 的 bill_of_materials）
      - 各步按 1:1 投料；有文献报道产率的步骤按产率放大，没有的按 100% 算，并在
        assumed_full_yield_steps 里列出是哪几步；
      - 试剂现在没有当量数据，mmol 和 mg 都是 null；
      - 不给 target_mg 时，按这次分析的规模档位折算（A12；data/scoring.toml 的 [scale]，mg 级是 10 mg）。
    怎么读：目标分子是阿司匹林（这里直接写了它的 SMILES），路线 r1 只有一步：水杨酸 + 乙酸酐，浓硫酸催化。
    目标产量取 180.16 mg，正好是 1 mmol 阿司匹林，这样期望的数字就是各个原料的分子量。
    """
    analysis = start_analysis(client, account, smiles="CC(=O)Oc1ccccc1C(=O)O")
    url = f"/api/analyses/{analysis['id']}/reports/materials"
    bom = client.get(url, headers=account["headers"], params={"route_id": "r1", "target_mg": 180.16}).json()
    rows = {row["name"]: row for row in bom["rows"]}
    # 180.16 mg 阿司匹林 = 1 mmol；没有产率记录按 100% 算，所以水杨酸和乙酸酐各 1 mmol
    # 毫克数 = 毫摩尔数 × 分子量：水杨酸的分子量是 138.12，乙酸酐是 102.09（RDKit 算的，保留两位小数）
    assert rows["水杨酸"]["mmol"] == 1.0 and rows["水杨酸"]["mg"] == 138.12
    # steps 是这个物料用在第几步
    assert rows["乙酸酐"]["mg"] == 102.09 and rows["乙酸酐"]["steps"] == [1]
    # 硫酸是试剂：列出来，但没有用量
    assert rows["硫酸"]["role"] == "试剂" and rows["硫酸"]["mg"] is None
    # 测试的先例库是空的，这一步没有文献产率，是按 100% 算的
    assert bom["assumed_full_yield_steps"] == [1]
    default = client.get(url, headers=account["headers"], params={"route_id": "r1"}).json()
    assert default["target_mg"] == 10  # mg 级默认按 10 mg 折算


def test_recording_an_experiment_marks_the_step_as_done(client, account, ctx):
    """G1–G4 回填实验结果：保存以后这一步显示“已做”；授权默认不勾、可以撤回，只有授权了的记录进入校准数据。

    保证的行为：
      - POST …/steps/{n}/experiments（api/reports.py 的 record_experiment，保存在 m14_backfill.py 的
        save_record）返回 201 和这条记录；label 是路线树上显示的文字（m14_backfill.py 的 experiment_label）；
      - 结果挂到分析结果里的这一步上（apply_to_result）：证据卡和路线树上都看得到（R12）；
      - 产率超过 100 被拒绝（api/schemas.py 的 ExperimentCreate）；
      - 接口返回的时间都带时区（syntherisk/models.py 的 iso）；
      - “授权用于改进模型”（G3）默认是 false；PATCH /api/experiments/{id}（update_experiment）可以补上或撤回；
      - 校准用的数据（m14_backfill.py 的 consented_dataset）只包含现在授权着的记录，而且是匿名的：
        只有反应和结构化的结果，没有备注、课题组和人员；run_calibration 在它上面做统计。
    怎么读：给格氏路线第 4 步记一次成功（产率 78%），第 2 步记一次失败；看证据卡、路线树、记录列表；
    然后在“没授权 → 授权 → 撤回”三个时刻各取一次校准数据。校准数据没有接口，
    测试直接开数据库会话调 m14_backfill 的函数（命令行的 calibrate、export-consented 调的也是它们）。
    为什么重要：m14_backfill.py 的文件头把这几条列为回填数据的规矩——授权默认不勾、可以随时撤回，
    导出的数据不带课题组、人员和自由填写的备注。
    """
    analysis = start_analysis(client, account)
    url = f"/api/analyses/{analysis['id']}/routes/r1/steps/4/experiments"
    # G1 结果、G2 产率和条件改动、备注。这里没有带 consent_for_training
    record = client.post(url, headers=account["headers"], json={
        "outcome": "success", "actual_yield": 78, "condition_changes": "乙醚换成 2-甲基四氢呋喃",
        "notes": "冬天湿度低，格氏试剂引发更快；干冰要现买",
    })
    assert record.status_code == 201, record.text
    body = record.json()
    assert body["label"] == "已做 · 实际 78%" and body["consent_for_training"] is False  # 授权默认不勾

    # 证据卡上这一步有了“已做”的标记
    card = client.get(f"/api/analyses/{analysis['id']}/routes/r1/steps/4", headers=account["headers"]).json()
    assert card["experiment"]["label"] == "已做 · 实际 78%"
    # 路线树：根是目标分子，它下面的那个反应节点（children[0]）就是最后一步，也就是第 4 步
    tree = client.get(f"/api/analyses/{analysis['id']}/routes/r1", headers=account["headers"]).json()["tree"]
    assert tree["children"][0]["experiment"]["actual_yield"] == 78

    # 给第 2 步记一次失败：结果是失败时，文字固定是“已做 · 失败”
    failed = client.post(f"/api/analyses/{analysis['id']}/routes/r1/steps/2/experiments", headers=account["headers"],
                         json={"outcome": "failure"}).json()
    assert failed["label"] == "已做 · 失败"
    # 产率是百分数，只能在 0 到 100 之间：120 通不过请求体的检查
    assert client.post(url, headers=account["headers"], json={"outcome": "success", "actual_yield": 120}).status_code == 422

    # 这次分析下的全部回填记录：第 2 步和第 4 步各一条
    listed = client.get(f"/api/analyses/{analysis['id']}/experiments", headers=account["headers"]).json()
    assert {item["step_index"] for item in listed} == {2, 4}
    # 时间总是带时区：不带的话，浏览器会把它当成本地时间
    assert all(item[field].endswith("+00:00") for item in listed for field in ("created_at", "updated_at"))

    # 没授权的记录不进入校准数据
    with ctx.session_factory() as session:
        assert m14_backfill.consented_dataset(session) == []
    # 补上授权：只带这一个字段，别的内容（比如备注）保持原样
    granted = client.patch(f"/api/experiments/{body['id']}", headers=account["headers"], json={"consent_for_training": True}).json()
    assert granted["consent_for_training"] is True and granted["notes"].startswith("冬天")
    with ctx.session_factory() as session:
        # 现在有一条授权了的记录。左边的写法要求右边的列表正好一个元素
        (row,) = m14_backfill.consented_dataset(session)
        # 匿名：只有反应和结构化的结果，没有备注、组织和人员
        assert set(row) == {"reaction_smiles", "outcome", "actual_yield", "step_score", "dimensions", "evidence_level",
                            "reported_yield_median", "operation_tags"}
        # operation_tags 是保存记录时这一步命中的操作规则：格氏羧基化命中“无水无氧操作”那一条
        assert row["actual_yield"] == 78 and row["operation_tags"] == ["anhydrous_organometallic"]
        # 校准统计：用了 1 条记录，结果是 1 次成功
        stats = m14_backfill.run_calibration(session).stats
        assert stats["n_records"] == 1 and stats["outcomes"] == {"success": 1}

    # 授权可以随时撤回
    client.patch(f"/api/experiments/{body['id']}", headers=account["headers"], json={"consent_for_training": False})
    with ctx.session_factory() as session:
        assert m14_backfill.consented_dataset(session) == []


def test_experiment_attachments_are_stored_encrypted(client, account, ctx):
    """G2 实验附件（TLC、NMR 的照片或谱图）：加密存盘，能原样下载；只收图片和 PDF；别的课题组下载不了。

    保证的行为：
      - POST /api/experiments/{id}/attachments（api/reports.py 的 add_attachment）返回 201 和下载地址；
        下载回来的内容和上传的一字节不差（download_attachment）；
      - 文件由 syntherisk/storage.py 的 save_file 加密以后写到 var/files 下，磁盘上看不到原文；
      - 类型按上传时声明的 Content-Type 判断，只接受图片（image/…）和 PDF，其余返回 415；
      - 附件按课题组隔离。
    怎么读：先给第 1 步记一条实验结果（附件要挂在一条记录上），上传一张假的 PNG 图片，下载回来比对，
    再直接去看磁盘上的文件；然后上传一个 exe；最后用别的课题组的账号去下载。
    """
    analysis = start_analysis(client, account)
    record = client.post(f"/api/analyses/{analysis['id']}/routes/r1/steps/1/experiments", headers=account["headers"],
                         json={"outcome": "partial", "actual_yield": 40}).json()
    # 一段假的图片内容：开头 8 个字节是 PNG 文件固定的文件头，后面是一段重复 50 遍的文字，方便在磁盘上搜
    picture = b"\x89PNG\r\n\x1a\n" + b"tlc-plate-photo" * 50
    # files 的值是 (文件名, 内容, Content-Type)。文件名里有中文和空格
    uploaded = client.post(f"/api/experiments/{record['id']}/attachments", headers=account["headers"],
                           files={"file": ("TLC 板.png", picture, "image/png")})
    assert uploaded.status_code == 201, uploaded.text
    assert client.get(uploaded.json()["download_url"], headers=account["headers"]).content == picture
    # ctx.settings.files_dir 是这个测试的 var/files 文件夹；把里面每个文件的字节都读出来。
    # 文件夹里确实有文件，而且哪个文件里都找不到那段文字：存的是密文
    on_disk = [path.read_bytes() for path in ctx.settings.files_dir.iterdir()]
    assert on_disk and all(b"tlc-plate-photo" not in data for data in on_disk)

    # 一个声明为可执行文件类型的上传：415（不支持的类型）
    rejected = client.post(f"/api/experiments/{record['id']}/attachments", headers=account["headers"],
                           files={"file": ("run.exe", b"MZ", "application/x-msdownload")})
    assert rejected.status_code == 415
    other = register(client, email="other@example.com", org="隔壁课题组")
    assert client.get(uploaded.json()["download_url"], headers=other["headers"]).status_code == 404
