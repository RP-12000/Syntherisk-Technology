"""E1 导入 Excel / CSV 库存，E2 确认匹配。

名称、CAS 号、SMILES 任一列都能用来识别（M1 转结构，M4 去重匹配）。
导入结果分三栏：已匹配（直接进库存）、需确认（给出猜测，等用户确认或改结构）、未识别。

在整个流程里的位置
    属于 PPT 的 M1（把名称、CAS 号变成结构）加 M4（库存）。对应库存页的两个按钮，接口都在 api/inventory.py：
        E1 导入 Excel / CSV   POST /api/inventory/imports
                                  → parse_table、create_import，然后小表格当场 process_import，大表格进队列
                              GET  /api/inventory/imports/{id}
                                  → import_view（看进度和三栏结果，R10）
        E2 确认匹配           POST /api/inventory/imports/{id}/rows/{row_id}
                                  → confirm_row
    大表格（超过 SYNC_ROW_LIMIT 行）不在请求里处理，由任务队列（taskqueue.py）调用 run_import 在后台跑。

一次导入的经过
    1. parse_table     读文件、认表头，得到一行一个字典（这一步不碰数据库）
    2. create_import   建一条导入记录（inventory_imports 表）和每行一条待处理记录（inventory_import_rows 表）
    3. process_import  逐行识别：能确定结构的直接写进库存（inventory_items 表），其余的留给用户
    4. confirm_row     用户对“需确认”“未识别”的行逐条处理：接受猜测、自己给结构，或者忽略

每一行所在的“栏”（InventoryImportRow.bucket）
    pending             还没处理
    matched             识别出来了，已经进库存
    needs_confirmation  有一个猜测，但不够确定（名称是别名、CAS 号校验位不对、是混合物……），等用户确认
    unrecognized        名称、CAS 号、SMILES 都没能识别
    confirmed           用户确认或者改了结构以后进了库存（返回给界面时和 matched 合在一起算“已匹配”）
    ignored             用户选了忽略

依赖
    modules/m01_resolver.py 的 NameResolver（即 ctx.resolver：先查内置试剂表，再查 PubChem、OPSIN，可能联网）、
    modules/m04_availability.py 的 InventoryService（写库存，按结构去重）、modules/m01_structure.py、
    external.py 的 is_cas_number、reference.py 的 normalize_name。Excel 用 openpyxl 读。
    库存条目和导入行的内容都是加密存的（models.py 里的 EncryptedJSON 列）。

本文件里有什么
    MAX_ROWS、SYNC_ROW_LIMIT  一次最多导入多少行；多少行以内当场处理
    _HEADERS、_HEADER_LOOKUP  表头别名表
    TableFormatError          表格读不了时抛的错误
    _cell                     把单元格的值变成文字
    _read_rows                把上传的文件读成“行的列表”
    parse_table               认表头，整理出每一行的各个字段
    create_import             建导入记录
    process_import            逐行识别并写库存
    run_import                任务队列的入口
    confirm_row               E2：处理一条不确定的行
    import_view               R10：一次导入的进度和结果
    row_view                  一行的显示内容
"""

from __future__ import annotations

import csv
import io
from typing import Any

from sqlalchemy import func, select
from sqlalchemy.orm import Session

from syntherisk import models
from syntherisk.external import is_cas_number
from syntherisk.modules import m01_structure as chem
from syntherisk.modules.m04_availability import InventoryService
from syntherisk.reference import normalize_name

# 一次导入最多多少行数据（空行不算）。超过就整个拒绝，请用户拆成几个文件
MAX_ROWS = 5000
SYNC_ROW_LIMIT = 200  # 不超过这么多行就在请求里直接处理完，更大的表放到后台

# 表头别名：左边是后端用的字段名，右边是表格里可以出现的列名。比较之前两边都先经过 normalize_name
# （转成小写，去掉空格、连字符、下划线、括号、逗号顿号），所以 “CAS No.”“Cas_No.”“cas no.” 都对得上 "casno."
_HEADERS = {
    # 名称（中文名、英文名、商品名都行）
    "name": ["名称", "品名", "化学品名称", "试剂名称", "试剂", "中文名", "英文名", "name", "chemical", "chemicalname",
             "compound", "productname", "description"],
    # CAS 号
    "cas": ["cas", "cas号", "casno", "casno.", "casnumber", "cas#", "casrn"],
    # 结构（SMILES）
    "smiles": ["smiles", "结构", "structure"],
    # 数量或规格，原样当文字存（“500 mL”“25 g”），不做换算
    "quantity": ["数量", "库存量", "存量", "规格", "包装", "quantity", "amount", "qty", "size"],
    # 存放位置
    "location": ["位置", "存放位置", "存放地点", "柜号", "location", "storage", "place"],
    # 备注
    "note": ["备注", "说明", "note", "notes", "remark", "remarks", "comment"],
}
# 反过来查的表：{规范化以后的列名: 字段名}，认表头时用
_HEADER_LOOKUP = {normalize_name(alias): field for field, aliases in _HEADERS.items() for alias in aliases}


class TableFormatError(ValueError):
    """表格读不了，或者里面找不到能用的列。

    报错信息是给用户看的中文。api/inventory.py 的 import_inventory 接住它，回 422 并把这句话原样返回。
    """


# ---------------------------------------------------------------- 读表格（不碰数据库）


def _cell(value: Any) -> str:
    """把一个单元格的值变成去掉首尾空白的文字；空单元格（None）变成空字符串。

    Excel 的单元格可能是数字、日期等类型，这里一律转成文字。
    谁在用：
        本文件的 _read_rows。
    """
    return "" if value is None else str(value).strip()


def _read_rows(filename: str, content: bytes) -> list[list[str]]:
    """把上传的文件读成“行的列表”，每行是一个单元格文字的列表。按文件名的扩展名决定怎么读。

    参数：
        filename  上传时的文件名，只用来看扩展名（不分大小写）。
        content   文件的全部内容（字节）。
    返回：
        [[第 1 行的各个单元格], [第 2 行的……], ...]。表头行也在里面，原样返回，这里不做任何识别。
        - .xlsx / .xlsm：只读第一个工作表；有公式的单元格取上次保存时算好的值。
        - .csv / .tsv / .txt：先按 UTF-8 解码，不行再按 GB18030（GBK 的超集；中文 Windows 上的 Excel
          另存 CSV 时默认用 GBK）；分隔符看开头 2000 个字符里制表符多还是逗号多。
    抛出：
        TableFormatError  Excel 文件打不开；文字两种编码都解不出来；扩展名不是上面这几种。
    谁在用：
        本文件的 parse_table。
    """
    lowered = filename.lower()
    if lowered.endswith((".xlsx", ".xlsm")):
        # 读 Excel 时才导入 openpyxl
        from openpyxl import load_workbook

        try:
            # read_only=True：流式读取，大文件也不占多少内存；data_only=True：公式取算好的值，不取公式本身
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception as exc:  # noqa: BLE001
            raise TableFormatError("这个 Excel 文件打不开") from exc
        sheet = workbook.worksheets[0]
        # values_only=True：直接给单元格的值，不给单元格对象
        return [[_cell(v) for v in row] for row in sheet.iter_rows(values_only=True)]
    if lowered.endswith((".csv", ".tsv", ".txt")):
        # for … else：循环里一次 break 都没发生（两种编码都失败）时，才会走到 else
        for encoding in ("utf-8-sig", "gb18030"):
            try:
                text = content.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        else:
            raise TableFormatError("这个文件的文字编码认不出来，请另存为 UTF-8 的 CSV")
        sample = text[:2000]
        delimiter = "\t" if sample.count("\t") > sample.count(",") else ","
        return [[_cell(v) for v in row] for row in csv.reader(io.StringIO(text), delimiter=delimiter)]
    raise TableFormatError("只支持 .xlsx 和 .csv 文件")


def parse_table(filename: str, content: bytes) -> list[dict[str, Any]]:
    """读表格，认出表头，返回每一行的 {name, cas, smiles, quantity, location, note, row_number}。

    参数：
        filename  上传时的文件名（看扩展名）。
        content   文件内容（字节）。
    返回：
        数据行的列表，每行一个字典：
            name、cas、smiles、quantity、location、note  表格里对应列的文字；没有这一列或者是空的就是 ""
            row_number                                  这一行在文件里是第几行（从 1 数起，表头也算一行）
        名称、CAS 号、SMILES 三样都空的行（空行）不在里面。
    规则：
        - 表头不一定在第一行：在前 10 行里找第一个含有“名称 / CAS / SMILES”任一列名的行当表头，
          它上面的行（标题、说明）都忽略。
        - 几列对上同一个字段时（比如既有“中文名”又有“英文名”），每行取最靠左的那个非空值。
        - 名称列里填的其实是 CAS 号（格式和校验位都对）、而 CAS 列是空的：挪到 cas 字段，name 清空。
    副作用：
        没有，不碰数据库。
    抛出：
        TableFormatError  文件读不了；前 10 行里找不到表头；数据行超过 MAX_ROWS；一行数据都没有。
    谁在用：
        api/inventory.py 的 import_inventory。
    """
    rows = _read_rows(filename, content)
    header_index, mapping = None, {}
    for index, row in enumerate(rows[:10]):
        # found：这一行里认得出来的列，{第几列: 字段名}
        found = {position: _HEADER_LOOKUP[normalize_name(cell)] for position, cell in enumerate(row)
                 if normalize_name(cell) in _HEADER_LOOKUP}
        # 至少要有一列能用来识别物质（名称、CAS、SMILES）才算表头；只有“数量”“位置”这类列的不算
        if any(field in found.values() for field in ("name", "cas", "smiles")):
            header_index, mapping = index, found
            break
    if header_index is None:
        raise TableFormatError("没找到表头：第一行需要有“名称”“CAS”或“SMILES”中的至少一列")

    parsed: list[dict[str, Any]] = []
    # enumerate 的第二个参数是起始编号：表头是文件的第 header_index + 1 行，它下面一行就是第 header_index + 2 行
    for number, row in enumerate(rows[header_index + 1 :], header_index + 2):
        record = {field: "" for field in _HEADERS}
        for position, field in mapping.items():
            # 这一行够长、而且这个字段还没填过，才填：先到先得，所以取到的是最靠左的非空值
            if position < len(row) and not record[field]:
                record[field] = row[position]
        # 有人把 CAS 号填在名称列里
        if not record["cas"] and is_cas_number(record["name"]):
            record["cas"], record["name"] = record["name"], ""
        # 没有任何可以用来识别的内容：当空行跳过
        if not any(record[field] for field in ("name", "cas", "smiles")):
            continue
        record["row_number"] = number
        parsed.append(record)
        if len(parsed) > MAX_ROWS:
            raise TableFormatError(f"一次最多导入 {MAX_ROWS} 行，请拆成几个文件")
    if not parsed:
        raise TableFormatError("表格里没有数据行")
    return parsed


# ---------------------------------------------------------------- 导入：建记录、逐行识别、进库存


def create_import(session: Session, org_id: str, user_id: str, filename: str, rows: list[dict[str, Any]]) -> models.InventoryImport:
    """给一次导入建记录：一条导入记录，加上每个数据行一条“待处理”的行记录。

    参数：
        session   数据库会话。
        org_id    课题组编号；导入的库存归这个课题组。
        user_id   谁上传的；进库存的条目以他为所有者。
        filename  上传时的文件名，只是记下来显示。
        rows      parse_table 的返回值。
    返回：
        新建的 InventoryImport：status 是默认的 "queued"，total_rows 是行数，已经有 id。
    副作用：
        往 inventory_imports 表加一行、往 inventory_import_rows 表加 len(rows) 行。只 flush 不提交，
        提交由调用方做。每条行记录的 bucket 是 "pending"，payload（加密存）是
        {"raw": 表格里的六个字段, "proposal": None}。
    谁在用：
        api/inventory.py 的 import_inventory。
    """
    record = models.InventoryImport(org_id=org_id, user_id=user_id, filename=filename, total_rows=len(rows))
    session.add(record)
    # flush：把上面的 INSERT 发给数据库但不提交。这样 record.id 就有值了，下面的行记录要用它
    session.flush()
    for row in rows:
        session.add(
            models.InventoryImportRow(
                import_id=record.id, org_id=org_id, row_number=row["row_number"], bucket="pending",
                # raw 只留表格里的六个字段（行号单独存在 row_number 列里）；proposal 等识别以后再填
                payload={"raw": {k: row[k] for k in _HEADERS}, "proposal": None},
            )
        )
    session.flush()
    return record


def process_import(ctx, session: Session, record: models.InventoryImport) -> None:
    """把一次导入里还没处理的行逐行识别一遍，能确定的直接写进库存。

    参数：
        ctx      运行时上下文；用 ctx.resolver 把名称、CAS 号变成结构。
        session  数据库会话。这个函数自己会提交，而且不止一次。
        record   create_import 建的导入记录。
    每一行的处理：
        交给 NameResolver.resolve_row，结论（matched / needs_confirmation / unrecognized）写进行记录的 bucket，
        识别的详情写进 payload["proposal"]（status、name、cas、smiles、reason、source、notes）。
        结论是 matched 的，立刻用 InventoryService.upsert 写进库存，并把库存条目的编号记在行记录的 item_id 上；
        库存里已经有同一个结构时是更新那一条（数量、位置、备注、CAS 号），不会多出一条。
        只处理 bucket 还是 "pending" 的行，所以重跑一遍时，处理过的行不会再处理。
    副作用：
        - 写数据库：导入记录的 status（先 "running"，最后 "done"）和 processed_rows，各行的 bucket、
          payload、item_id，以及 inventory_items 表。开始时提交一次，之后每处理 20 行提交一次，
          最后再提交一次：别的请求（GET /api/inventory/imports/{id}）因此看得到进度，
          中途出错也只丢最近不到 20 行的结果。
        - 可能联网：内置试剂表里查不到的名称和 CAS 号会去问 PubChem、OPSIN（开着的话）。
    抛出：
        不主动抛错。数据库出问题这类意外会原样抛出去，这时导入记录停在 "running"；
        从任务队列进来时由 run_import 接住并记成 "failed"。
    谁在用：
        api/inventory.py 的 import_inventory（不超过 SYNC_ROW_LIMIT 行时在请求里直接调用）；
        本文件的 run_import（大表格，在后台）。
    """
    resolver = ctx.resolver
    # 库存服务绑定到“这个课题组、这个上传者”：写进去的条目归他所有
    inventory = InventoryService(session, record.org_id, record.user_id)
    record.status = "running"
    session.commit()
    # session.scalars(select(模型)…).all()：查出符合条件的行，直接得到模型对象的列表。
    # 只取还没处理的行，按表格里的行号排好
    pending = session.scalars(
        select(models.InventoryImportRow)
        .where(models.InventoryImportRow.import_id == record.id, models.InventoryImportRow.bucket == "pending")
        .order_by(models.InventoryImportRow.row_number)
    ).all()
    for count, row in enumerate(pending, 1):
        raw = row.payload["raw"]
        # 识别这一行：SMILES 优先，其次 CAS 号，最后名称（细节见 m01_resolver.NameResolver.resolve_row）
        resolution = resolver.resolve_row(name=raw["name"], cas=raw["cas"], smiles=raw["smiles"])
        row.bucket = resolution.status
        # payload 是加密列：要整个重新赋值才会保存，只改里面的某个键不会被存下来
        row.payload = {"raw": raw, "proposal": resolution.to_dict()}
        if resolution.status == "matched":
            # upsert 返回 (库存条目, 是不是新建的)；第二项这里用不到，变量名前面的下划线就是这个意思
            item, _created = inventory.upsert(
                name=resolution.name, cas=resolution.cas, smiles=resolution.smiles, quantity=raw["quantity"],
                location=raw["location"], note=raw["note"], source="import", import_id=record.id,
            )
            row.item_id = item.id
        record.processed_rows = count
        # 每 20 行提交一次，进度才看得见
        if count % 20 == 0:
            session.commit()
    record.status = "done"
    session.commit()


def run_import(ctx, import_id: str) -> None:
    """任务队列的入口（大文件在后台跑）。

    参数：
        ctx        运行时上下文。
        import_id  导入记录的编号（入队时放在任务内容里的 import_id）。
    做什么：
        自己开一个数据库会话，取出导入记录，调用 process_import。
        记录不存在、或者已经是 "done"，就什么都不做（同一个任务被重复投递也不会重复导入）。
    副作用：
        见 process_import。出错时回滚没提交的部分，把导入记录的 status 改成 "failed"、
        error 写成 “错误类型: 错误信息”（最多 1000 个字符）并提交，然后把异常继续往外抛，
        由 taskqueue.run_job 记进日志。
    谁在用：
        taskqueue.py 的 job_handlers 把任务类型 "inventory.import" 对到这个函数；
        api/inventory.py 的 import_inventory 在表格超过 SYNC_ROW_LIMIT 行时入队；tests/test_app_service.py。
    """
    with ctx.session_factory() as session:
        # session.get(模型, 主键)：按主键取一行，没有就是 None
        record = session.get(models.InventoryImport, import_id)
        if record is None or record.status == "done":
            return
        try:
            process_import(ctx, session, record)
        except Exception as exc:  # noqa: BLE001
            # 先回滚，把出错前还没提交的改动丢掉，这个会话才能接着用来记下失败原因
            session.rollback()
            record.status = "failed"
            record.error = f"{type(exc).__name__}: {exc}"[:1000]
            session.commit()
            raise


# ---------------------------------------------------------------- E2 确认，和给界面看的结果（R10）


def confirm_row(
    session: Session, row: models.InventoryImportRow, user_id: str, *, action: str, smiles: str | None = None, name: str | None = None
) -> models.InventoryImportRow:
    """E2：confirm 接受系统的猜测；set_structure 用用户给的结构；ignore 这一行不要了。

    参数：
        session  数据库会话。
        row      要处理的那一行。接口已经检查过它属于这次导入、而且还在“需确认”或“未识别”一栏。
        user_id  做这个操作的人；新建的库存条目以他为所有者。
        action   "confirm" / "set_structure" / "ignore"（接口的请求体只接受这三个值）。
        smiles   action 是 set_structure 时用户给的结构；其他时候不用。
        name     用户改的名字，可以不给。不给时依次用表格里的名称、系统猜的名称，都没有就叫“未命名”。
    返回：
        传进来的同一个 row（已经改好）。
    三种操作：
        ignore         bucket 改成 "ignored"，不进库存。
        confirm        用 proposal 里猜的结构进库存。“未识别”的行没有猜测，也可以 confirm：
                       这时不带结构，只按名称记进库存（混合物、自制的东西就是这样）。
        set_structure  把用户给的 SMILES 转成规范写法，用它进库存。
        后两种：数量、位置、备注取表格里的；CAS 号先用 proposal 里的，没有再用表格里的。
        进库存用 InventoryService.upsert，库存里已有同一个结构时更新那一条。
        之后 bucket 改成 "confirmed"，item_id 记下库存条目的编号。
    副作用：
        改 row；写 inventory_items 表（upsert 里新建条目时会 flush）。不提交，提交由调用方做。
    抛出：
        ValueError  set_structure 给的 SMILES 读不出来；或者 action 不是上面三个值。接口把它转成 422。
    谁在用：
        api/inventory.py 的 resolve_import_row。
    """
    if action == "ignore":
        row.bucket = "ignored"
        return row
    # raw 是表格里的原样；proposal 是识别时留下的猜测，没有的话当成空字典
    raw, proposal = row.payload["raw"], row.payload.get("proposal") or {}
    if action == "set_structure":
        # canonical_smiles：解析并转成规范写法，读不出来返回 None
        canonical = chem.canonical_smiles(smiles or "")
        if canonical is None:
            raise ValueError("给的结构读不出来")
        final_smiles = canonical
    elif action == "confirm":
        final_smiles = proposal.get("smiles")  # 混合物可以没有结构，照样记进库存
    else:
        raise ValueError(f"不认识的操作：{action}")
    inventory = InventoryService(session, row.org_id, user_id)
    # 名字按这个先后取：用户这次给的 → 表格里的 → 系统猜的 → “未命名”。
    # CAS 号先用系统识别时得到的，没有再用表格里的。数量、位置、备注都取表格里的原样
    item, _created = inventory.upsert(
        name=name or raw["name"] or proposal.get("name") or "未命名",
        cas=proposal.get("cas") or raw["cas"],
        smiles=final_smiles,
        quantity=raw["quantity"],
        location=raw["location"],
        note=raw["note"],
        source="import",
        import_id=row.import_id,
    )
    row.item_id = item.id
    row.bucket = "confirmed"
    return row


def import_view(session: Session, record: models.InventoryImport) -> dict[str, Any]:
    """R10：三栏的计数，加上需确认和未识别的明细。

    参数：
        session  数据库会话（只读）。
        record   导入记录。
    返回：
        {
          "id":             导入编号,
          "filename":       上传时的文件名,
          "status":         "queued" 排队中 / "running" 处理中 / "done" 完成 / "failed" 出错,
          "error":          出错时的原因，否则是 None,
          "total_rows":     一共多少行数据,
          "processed_rows": 已经处理了多少行（大表格在后台处理时看进度用）,
          "counts": {
              "matched":            已匹配 = 自动识别进库存的（matched）+ 用户确认过的（confirmed）,
              "needs_confirmation": 需确认,
              "unrecognized":       未识别,
              "ignored":            用户忽略的,
              "pending":            还没处理的,
          },
          "rows":           需确认和未识别两栏的每一行（row_view 的结果），按表格里的行号排；
                            已匹配、已确认、已忽略的行不在里面,
        }
    谁在用：
        api/inventory.py 的 import_inventory、get_import、resolve_import_row。
    """
    # SELECT bucket, COUNT(*) … GROUP BY bucket：每一栏有多少行，再转成 {栏: 行数}
    counts = dict(
        session.execute(
            select(models.InventoryImportRow.bucket, func.count())
            .where(models.InventoryImportRow.import_id == record.id)
            .group_by(models.InventoryImportRow.bucket)
        ).all()
    )
    # 明细只取还要用户处理的两栏
    rows = session.scalars(
        select(models.InventoryImportRow)
        .where(
            models.InventoryImportRow.import_id == record.id,
            models.InventoryImportRow.bucket.in_(("needs_confirmation", "unrecognized")),
        )
        .order_by(models.InventoryImportRow.row_number)
    ).all()
    return {
        "id": record.id,
        "filename": record.filename,
        "status": record.status,
        "error": record.error,
        "total_rows": record.total_rows,
        "processed_rows": record.processed_rows,
        "counts": {
            # counts.get(栏, 0)：这一栏一行都没有时，查询结果里就没有这个键，按 0 算
            "matched": counts.get("matched", 0) + counts.get("confirmed", 0),
            "needs_confirmation": counts.get("needs_confirmation", 0),
            "unrecognized": counts.get("unrecognized", 0),
            "ignored": counts.get("ignored", 0),
            "pending": counts.get("pending", 0),
        },
        "rows": [row_view(row) for row in rows],
    }


def row_view(row: models.InventoryImportRow) -> dict[str, Any]:
    """导入结果里一行的显示内容。

    返回：
        {
          "id":         行记录的编号（E2 接口路径里的 {row_id}）,
          "row_number": 在表格里是第几行,
          "bucket":     所在的栏，取值见文件头,
          "raw":        表格里的原样：name、cas、smiles、quantity、location、note,
          "proposal":   系统猜的结果 {name, cas, smiles, source}，source 是怎么查到的
                        （比如 “名称 · 内置试剂表”）；没有猜出结构时是 None,
          "reason":     为什么要人确认、或者为什么没识别出来（一句话，可能是空字符串）,
          "item_id":    已经进库存的话是库存条目的编号，否则是 None,
        }
    谁在用：
        本文件的 import_view；api/inventory.py 的 resolve_import_row。
    """
    proposal = row.payload.get("proposal") or {}
    return {
        "id": row.id,
        "row_number": row.row_number,
        "bucket": row.bucket,
        "raw": row.payload["raw"],
        # 只挑出四个键给界面；proposal 里没有 smiles（没猜出结构）时整个给 None
        "proposal": {k: proposal.get(k) for k in ("name", "cas", "smiles", "source")} if proposal.get("smiles") else None,
        "reason": proposal.get("reason", ""),
        "item_id": row.item_id,
    }
