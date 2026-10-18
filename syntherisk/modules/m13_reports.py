"""M13 报告生成：PDF、Word、SDF + RXN（F1–F3），以及对比表导出（D4）。

报告只用已经算好的结果和 M11 写好的说明，不重新分析，所以是秒级。
免责声明固定包含，不能取消。

对应关系
    PPT 模块：M13 报告生成。它不在流水线的八个阶段里，用户点“导出”时由 api/reports.py 调用。
    界面：
        F1 选择内容     sections：route / evidence_cards / literature / materials / safety（免责声明总是有）
        F2 选择格式     format：pdf / docx / sdf（sdf 实际给的是一个 zip：SDF 加每一步的 RXN）
        F3 生成并下载   POST /api/analyses/{id}/reports，之后 GET /api/reports/{id}/download
        D4 导出对比     POST /api/analyses/{id}/compare/export，xlsx 或 csv
        物料清单预览    GET /api/analyses/{id}/reports/materials?route_id=…&target_mg=…
        A12 规模        决定物料清单默认按多少毫克目标物折算（data/scoring.toml 的 [scale]：
                        mg 10、hundred_mg 100、g 1000）

每个函数出什么
    build_blocks()           不出文件。把报告内容整理成一串“块”（标题、段落、表格、图片……），和格式无关
    render_pdf()             PDF（ReportLab），输入是 build_blocks() 的块
    render_docx()            Word 的 .docx（python-docx），输入是同一串块
    export_structures_zip()  zip 压缩包：route.sdf（全部分子）、step_N.rxn（每一步的反应）、README.txt
    comparison_xlsx()        Excel 的 .xlsx（openpyxl）：“路线对比”和“权重”两张表
    comparison_csv()         CSV（带 BOM 的 UTF-8）：对比表
    bill_of_materials()      不出文件。算物料清单，返回字典（报告第四节和预览接口用）
    出文件的函数都只返回文件内容（bytes）。把文件加密存到 var/files、在 reports 表里记一行、
    提供下载，这些都在 api/reports.py 和 syntherisk/storage.py 里做。

数据从哪来
    分析结果 Analysis.result（molecules、reactions、routes 三张表）和 M10 写在路线上的分数、汇总；
    每一步的内容通过 syntherisk/cards.py 的 build_card() 取（和界面上的证据卡是同一份数据）；
    结构图和反应式图用 m01_structure 里的 RDKit 函数现画。
    本文件不读写数据库，不联网，不调大模型。

本文件里有什么
    SECTIONS、DISCLAIMER、STARS、Block   可选的部分、免责声明、星级的写法、“块”的类型
    _evidence()、_operation()、_fmt()     把证据等级、操作标签、数值变成表格里的文字
    default_filename()、safe_filename()   文件名
    bill_of_materials()                   物料清单
    build_blocks()                        报告内容（和格式无关）
    _PDF_FONT、_FONT_CANDIDATES、_pdf_font()   找一个 PDF 能用的中文字体
    render_pdf()                          块 → PDF
    render_docx()                         块 → Word
    export_structures_zip()               SDF + RXN 的压缩包
    _comparison_rows()                    对比表的表头和各行
    comparison_xlsx()、comparison_csv()   对比表 → Excel / CSV
"""

from __future__ import annotations

import csv
import io
import os
import zipfile
from datetime import date
from typing import Any
from xml.sax.saxutils import escape

from rdkit import Chem
from rdkit.Chem import AllChem

from syntherisk import i18n
from syntherisk.cards import build_card
from syntherisk.constraints import SCALE_LABELS
from syntherisk.modules import m01_structure as chem
from syntherisk.reference import ReferenceData
from syntherisk.routes import iter_mol_nodes, iter_reaction_nodes

# 报告可选的五个部分：编号 -> 中文名。编号就是 F1 提交的 sections 里的值
# （api/schemas.py 的 ReportCreate 里另有一份同样的编号，用来校验请求）。
# 这个字典本身现在没有被别的代码引用；报告里各节的标题是 build_blocks() 里直接写的。
SECTIONS = {
    "route": "推荐路线与路线树",
    "evidence_cards": "每一步的证据卡",
    "literature": "文献列表",
    "materials": "试剂与物料清单",
    "safety": "安全汇总：高危物质、特殊操作",
}
# 免责声明。PDF 和 Word 报告的最后一节、zip 包里的 README.txt 都会带上它，用户不能取消。
DISCLAIMER = (
    "本报告由 Syntherisk 根据公开数据和规则自动生成，仅供有资质的研究人员参考，"
    "不是实验操作规程，也不是安全建议。分数和证据等级用来辅助判断，不保证反应可行。"
    "动手之前，请核对原始文献、各试剂的安全技术说明书（SDS）和所在单位的安全规定。"
)
# 证据星级（0–3）在报告里的写法。星级来自 data/scoring.toml 的 [evidence.levels]：E1 三星……E4 零星。
STARS = {3: "★★★", 2: "★★☆", 1: "★☆☆", 0: "☆☆☆"}

# “块”：报告里的一段内容，是一个带 type 的字典。有哪些 type、各带什么字段，见 build_blocks() 的说明。
Block = dict[str, Any]


def _evidence(card: dict[str, Any]) -> str:
    """证据等级在报告里的写法：星级加等级名，例如“★★☆ 高相似先例”。

    没有评估时（先例库没导入，stars 是 None）只写等级的说明文字；连说明文字也没有，就写“未评估”。

    参数：
        card  证据卡；读 card["evidence"] 的 stars 和 label。
    返回：
        字符串。
    谁在用：
        本文件的 build_blocks()：路线表的“证据”一列、证据卡表的“证据等级”一行。
    """
    evidence = card["evidence"]
    if evidence["stars"] is None:
        return evidence["label"] or "未评估"
    return f"{STARS[evidence['stars']]} {evidence['label']}"


def _operation(card: dict[str, Any]) -> str:
    """操作标签形如“操作 中等 · 需无水”；放进表头已经是“操作”的格子里时去掉前缀。

    参数：
        card  证据卡；读 card["conditions"]["operation_label"]。
    返回：
        去掉开头“操作 ”以后的文字，例如“中等 · 需无水”；没有操作标签时返回“—”。
        removeprefix() 只在字符串以这个前缀开头时才去掉它，否则原样返回。
    谁在用：
        本文件的 build_blocks()。
    """
    return (card["conditions"].get("operation_label") or "—").removeprefix("操作 ")


def _fmt(value: Any, suffix: str = "") -> str:
    """把一个值变成表格里显示的文字。

    None 显示成“—”（不带后缀）；浮点数用 :g 去掉多余的小数位（75.0 显示成 75）；别的值原样转成文字。

    参数：
        value   要显示的值。
        suffix  接在后面的单位，比如 "%"。
    谁在用：
        本文件的 build_blocks()。
    """
    if value is None:
        return "—"
    if isinstance(value, float):
        return f"{value:g}{suffix}"
    return f"{value}{suffix}"


def default_filename(target_name: str, route_name: str, extension: str) -> str:
    """没有指定文件名时用的默认文件名：“目标物_路线名_日期.扩展名”。

    例如“布洛芬_格氏羧基化_2026-10-11.pdf”。日期是服务器当天的日期。

    参数：
        target_name  目标分子的名字（没有名字时调用方传的是分子式）。
        route_name   路线名；导出对比表时调用方传的是“路线对比”。
        extension    扩展名，不带点：pdf / docx / zip / xlsx / csv。
    返回：
        经过 safe_filename() 处理的文件名。
    谁在用：
        api/reports.py 的 create_report()（F3）和 export_comparison()（D4）。
    """
    # 文件名是当场起的，不走“先存中文、返回时再换”的对照（下载时的响应头不经过那一步）：
    # 这次请求要英文时，目标物和路线的名字有英文的就直接用英文。中文请求时两个名字原样
    if i18n.current_lang() == "en":
        target_name, route_name = i18n.en_of(target_name), i18n.en_of(route_name)
    stem = f"{target_name}_{route_name}_{date.today().isoformat()}"
    return safe_filename(f"{stem}.{extension}")


def safe_filename(name: str) -> str:
    """把一个名字变成可以放心用作文件名的样子。

    做三件事：把 Windows 文件名里不允许的字符（反斜杠、斜杠、冒号、星号、问号、双引号、尖括号、竖线）
    和回车、换行、制表符都换成下划线；去掉首尾的空格和点；最多留 150 个字符。
    处理完是空的，就叫 "report"。

    参数：
        name  原始的名字：用户填的报告文件名、上传的附件名，或者 default_filename() 拼出来的名字。
    返回：
        处理后的文件名。超过 150 个字符时从后面截断，扩展名可能被截掉；
        create_report() 发现文件名不是以扩展名结尾时会再补上。
    谁在用：
        本文件的 default_filename()；api/reports.py 的 create_report()（用户自己填了文件名时）
        和 add_attachment()（G2 上传的附件名）。
    """
    # 逐个字符看：在那串“不允许的字符”里的换成下划线，其余原样；最后去掉首尾的空格和点
    cleaned = "".join("_" if ch in '\\/:*?"<>|\r\n\t' else ch for ch in name).strip(" .")
    return cleaned[:150] or "report"


# ---------------------------------------------------------------- 物料清单


def bill_of_materials(result: dict[str, Any], route: dict[str, Any], target_mg: float) -> dict[str, Any]:
    """按目标产量倒推每个起始原料的用量。

    假设：各步按 1:1 化学计量投料；有文献报道产率的步骤按报道值的中位数放大，没有的按 100% 算。
    试剂的当量现在没有数据，只列出名称。

    参数：
        result     分析结果（用 molecules 和 reactions 两张表）。
        route      路线；用它的 tree。路线树里分子节点和反应节点交替出现，根是目标分子，
                   没有 children 的分子节点是起始原料（结构见 syntherisk/routes.py）。
        target_mg  想得到多少毫克目标物。
    返回（字典）：
        target_mg                 同参数
        rows                      物料列表，先是起始原料，后是试剂。每项：
            name        名字；没有名字时用分子式
            role        “起始原料”或“试剂”
            steps       用在第几步（序号从小到大）
            mol_weight  分子量
            mmol        用量，毫摩尔，保留四位小数；试剂是 None
            mg          用量，毫克，保留两位小数；试剂是 None
            source      来源：库存 / 可购买 / 买不到；不知道时是“未知”。试剂不在目录里也写“未知”，
                        和 m10_scoring.step_materials() 的处理一致
        assumed_full_yield_steps  没有产率记录、按 100% 算的步骤序号（从小到大，不重复）
        note                      固定的一句说明：按 1:1 估算，没有计入过量和损耗
    算法：
        目标物的毫摩尔数 = target_mg ÷ 目标物的分子量。从目标分子沿路线树往回走：要做出 n 毫摩尔产物，
        这一步的每个反应物各需要 n ÷ (这一步的产率 ÷ 100) 毫摩尔。一直走到起始原料，
        同一个原料在几处用到的，需要量加在一起。
    例子（tests/test_reports_and_backfill.py）：
        180.16 mg 阿司匹林是 1 mmol；那一步没有产率记录，按 100% 算，
        所以水杨酸要 1 mmol（138.12 mg），乙酸酐要 1 mmol（102.09 mg）；硫酸是试剂，没有用量。
    注意：
        起始原料里带 ubiquitous 角色的（现在只有水）不列；试剂没有做这个过滤。
    副作用：没有。
    谁在用：
        本文件的 build_blocks()（报告第四节）；
        api/reports.py 的 preview_materials()（GET /api/analyses/{id}/reports/materials）。
    """
    molecules, reactions = result["molecules"], result["reactions"]
    target = molecules[route["tree"]["key14"]]
    # needs：起始原料的 key14 -> 一共需要多少毫摩尔
    needs: dict[str, float] = {}
    # used_in：起始原料的 key14 -> 用在哪几步（步骤序号的集合）
    used_in: dict[str, set[int]] = {}
    # assumed：没有产率记录、按 100% 算的步骤序号
    assumed: list[int] = []

    def walk(node: dict[str, Any], mmol: float, consumer: int | None) -> None:
        """从一个分子节点往回走，算出做 mmol 毫摩尔这个分子需要多少起始原料。

        参数：
            node      路线树里的一个分子节点。
            mmol      这个分子需要多少毫摩尔。
            consumer  用掉这个分子的那一步的序号；根节点（目标分子）没有，是 None。
        副作用：往外层的 needs、used_in、assumed 里记数。它调用自己来走完整棵树（递归）。
        """
        if not node["children"]:
            # 没有 children：这是起始原料。把需要量累加上去，并记下它用在哪一步
            needs[node["key14"]] = needs.get(node["key14"], 0.0) + mmol
            if consumer:
                used_in.setdefault(node["key14"], set()).add(consumer)
            return
        # 有 children：这个分子是由一步反应做出来的。children 里只有一个反应节点
        reaction_node = node["children"][0]
        reaction = reactions[reaction_node["step_key"]]
        # 这一步的产率：文献报道值的中位数（0–100）；没有记录时是 None
        reported = (reaction.get("yield") or {}).get("median")
        if not reported:
            assumed.append(reaction_node["step_index"])
        for child in reaction_node["children"]:
            # 每个反应物需要 mmol ÷ 产率（没有产率记录时按 100% 算，也就是除以 1）
            walk(child, mmol / ((reported or 100) / 100), reaction_node["step_index"])

    # 从目标分子开始：毫克 ÷ 分子量 = 毫摩尔
    walk(route["tree"], target_mg / target["mol_weight"], None)

    # 清单里的几个固定用语。i18n.tr 返回的就是中文，同时登记英文，接口返回这份清单时按请求的语言换；
    # 报告正文（build_blocks）用的也是这份清单，那里一直是中文
    role_starting = i18n.tr("起始原料", "Starting material")
    role_reagent = i18n.tr("试剂", "Reagent")
    unknown = i18n.tr("未知", "Unknown")
    # 起始原料：每样一行，毫摩尔 × 分子量 = 毫克
    rows = []
    for key, mmol in needs.items():
        molecule = molecules[key]
        availability = molecule.get("availability") or {}
        if availability.get("ubiquitous"):
            continue
        rows.append(
            {
                "name": molecule.get("name") or molecule["formula"],
                "role": role_starting,
                "steps": sorted(used_in.get(key, [])),
                "mol_weight": molecule["mol_weight"],
                "mmol": round(mmol, 4),
                "mg": round(mmol * molecule["mol_weight"], 2),
                "source": availability.get("label", unknown),
            }
        )
    # 试剂：只知道哪一步用了它，不知道用多少。iter_reaction_nodes() 按实验顺序给出路线树里的每个反应节点
    reagent_steps: dict[str, set[int]] = {}
    for reaction_node in iter_reaction_nodes(route["tree"]):
        for key in reactions[reaction_node["step_key"]]["reagents"]:
            reagent_steps.setdefault(key, set()).add(reaction_node["step_index"])
    for key, steps in reagent_steps.items():
        molecule = molecules[key]
        availability = molecule.get("availability") or {}
        rows.append(
            {
                "name": molecule.get("name") or molecule["formula"],
                "role": role_reagent,
                "steps": sorted(steps),
                "mol_weight": molecule["mol_weight"],
                "mmol": None,
                "mg": None,
                "source": availability.get("label") if availability.get("status") != "unavailable" else unknown,
            }
        )
    return {
        "target_mg": target_mg,
        "rows": rows,
        "assumed_full_yield_steps": sorted(set(assumed)),
        "note": i18n.tr("按 1:1 化学计量估算，没有计入过量和损耗；试剂用量请参照文献当量。",
                        "Estimated at 1:1 stoichiometry, with no allowance for excess or losses; "
                        "for reagent amounts, follow the equivalents given in the literature."),
    }


# ---------------------------------------------------------------- 报告内容（和格式无关）


def build_blocks(
    result: dict[str, Any],
    route: dict[str, Any],
    reference: ReferenceData,
    *,
    sections: list[str],
    scale: str,
    target_mg: float | None = None,
) -> list[Block]:
    """把一份报告的内容整理成一串“块”，和输出格式无关。PDF 和 Word 用的是同一串块。

    参数：
        result     分析结果。
        route      要出报告的路线（评过分的；带着 explanation 的话，推荐理由会放在第一节开头）。
        reference  参考数据；用 scoring（各规模对应的毫克数）。build_card() 还要用它取维度名和权重。
        sections   要包含哪几部分（F1）：route / evidence_cards / literature / materials / safety 里的若干个。
                   不在列表里的部分整个不出现。开头的标题和目标分子图、最后的免责声明总是有。
        scale      规模（A12）：mg / hundred_mg / g。显示在报告开头；没给 target_mg 时，物料清单按它折算。
        target_mg  物料清单按得到多少毫克目标物折算；None 时用 data/scoring.toml [scale] 里这个规模的数。
    返回：
        块的列表。每个块是一个字典，type 决定它是什么：
            title / h1 / h2 / p   报告标题、一级标题、二级标题、正文段落：{"type", "text"}
            caption / note        图题、小字说明：{"type", "text"}
            bullets               圆点列表：{"type", "items": [文字, ...]}
            image                 图片：{"type", "png": PNG 图片的字节（画不出来时是 None）,
                                        "width_mm": 宽度（毫米）, "ratio": 高 ÷ 宽}
            table                 表格：{"type", "header": [表头], "rows": [[单元格, ...], ...],
                                        "widths": [各列的相对宽度]}
            pagebreak             分页：{"type"}
    各部分的内容：
        开头（总是有）
            标题“<目标物> 合成路线报告”；一行：路线名、生成日期、规模；目标分子的结构图
        一、推荐路线（sections 里有 route）
            推荐理由（M11 写的 route["explanation"]；没有就没有这一段）
            一行汇总表：综合分、置信度和理由、步数、总产率、原料成本、最高危害、文献覆盖、库存覆盖
            本次约束（result["constraint_tags"] 的标签；没有就不写）
            路线表：每一步的反应名、得分、证据等级、操作
            每一步的反应式图，图上方是“第 N 步 · 反应名”
            最薄弱环节的说明和建议（有才写）
        二、每一步的证据卡（evidence_cards，另起一页）
            每一步：标题和得分；M11 写的这一步的说明（有才写）；
            一张“项目 / 内容”表：证据等级和置信度、文献依据、反应条件、操作、产率、毒性、原料、绿色化学；
            一张维度表：六个维度的权重、得分、依据
        三、文献列表（literature）
            每条先例一行：第几步、编号、类型、文献号、年份、相似度（同一反应写“同一反应”）、产率、链接；
            各步都没有先例时只写一句说明
        四、试剂与物料清单（materials）
            bill_of_materials() 的结果：物料、角色、用在第几步、分子量、用量（mmol 和 mg）、来源；
            表下面一行小字说明估算的假设
        五、安全汇总（safety）
            物质表：物质、用在第几步、GHS 危险性说明（最多五条）、管制情况。只列有危险性说明、
            在管制清单上、或者没有危害数据的物质
            特殊操作：各步命中的操作标签和理由（有才写）
            有危害分类取自内置试剂表、还没有复核时，加一行提醒
        免责声明（总是有）
            DISCLAIMER；分析结果里的每条数据说明（result["warnings"]）各一行小字
    章节的序号是写死的：只选一部分时序号不会重排（只选 route 和 materials，标题就是“一”和“四”）。
    副作用：没有。
    谁在用：
        api/reports.py 的 create_report()（F3）；返回值交给 render_pdf() 或 render_docx()。
    """
    scoring = reference.scoring
    target = result["target"]
    target_name = target.get("name") or target["formula"]
    # 每一步一张证据卡（和界面上 R6 的证据卡是同一个函数、同一份数据），后面各节都从卡里取内容
    cards = [build_card(result, route, step["index"], reference) for step in route["steps"]]
    summary = route["summary"]
    # ---- 开头：标题、一行基本信息、目标分子的结构图。
    # 图按 420×260 像素画，在报告里宽 70 毫米；ratio 是高 ÷ 宽，渲染时用它算出图的高度
    blocks: list[Block] = [
        {"type": "title", "text": f"{target_name} 合成路线报告"},
        {"type": "p", "text": f"路线：{route['name']}　·　生成日期：{date.today().isoformat()}　·　规模：{SCALE_LABELS[scale]}"},
        {"type": "image", "png": chem.mol_to_png(target["smiles"], 420, 260), "width_mm": 70, "ratio": 260 / 420},
    ]

    # ---- 一、推荐路线
    if "route" in sections:
        blocks.append({"type": "h1", "text": "一、推荐路线"})
        explanation = route.get("explanation")
        if explanation:
            blocks.append({"type": "p", "text": explanation["text"]})
        # 汇总表只有一行，数据来自 M10 写在路线上的 score、confidence_* 和 summary。
        # widths 是各列的相对宽度：置信度那一列带理由，文字长，所以给 2.6
        blocks.append({
            "type": "table",
            "header": ["综合分", "置信度", "步数", "总产率", "原料成本", "最高危害", "文献覆盖", "库存覆盖"],
            "rows": [[
                _fmt(route["score"]),
                f"{route['confidence_label']}（{route['confidence_reason']}）",
                str(summary["n_steps"]),
                _fmt(summary["overall_yield"], "%") if summary["overall_yield"] is not None else "无完整产率记录",
                summary["material_cost_label"],
                summary["max_hazard"]["label"] if summary["max_hazard"] else "—",
                _fmt(summary["literature_coverage"], "%") if summary["literature_coverage"] is not None else "未评估",
                _fmt(summary["inventory_coverage"], "%"),
            ]],
            "widths": [1, 2.6, 0.8, 1.4, 1.1, 1.3, 1.1, 1.1],
        })
        # 这次分析用的约束（R2 的约束标签），拼成一句
        tags = "、".join(tag["label"] for tag in result.get("constraint_tags") or [])
        if tags:
            blocks.append({"type": "p", "text": f"本次约束：{tags}"})
        blocks.append({"type": "h2", "text": "路线（从第 1 步做到最后一步）"})
        blocks.append({
            "type": "table",
            "header": ["步骤", "反应", "得分", "证据", "操作"],
            "rows": [[f"第 {c['step_index']} 步", c["name"], _fmt(c["score"]), _evidence(c), _operation(c)] for c in cards],
            "widths": [0.9, 2.2, 0.8, 2.4, 2.4],
        })
        # 每一步的反应式图：反应物 → 产物，按 720×180 像素画，在报告里宽 165 毫米。
        # reaction_to_png() 画不出反应式时会退回只画产物；连产物也画不出来（返回 None）就跳过这一步
        for card in cards:
            png = chem.reaction_to_png([m["smiles"] for m in card["reactants"]], card["product"]["smiles"], 720, 180)
            if png:
                blocks.append({"type": "caption", "text": card["title"]})
                blocks.append({"type": "image", "png": png, "width_mm": 165, "ratio": 180 / 720})
        weakest = route.get("weakest")
        if weakest:
            blocks.append({"type": "p", "text": f"最薄弱环节：第 {weakest['step_index']} 步 {weakest['name']}（{_fmt(weakest['score'])} 分）。"
                                                f"{weakest['text']}{weakest['suggestion']}。"})

    # ---- 二、每一步的证据卡（另起一页）
    if "evidence_cards" in sections:
        blocks.append({"type": "pagebreak"})
        blocks.append({"type": "h1", "text": "二、每一步的证据卡"})
        for card in cards:
            blocks.append({"type": "h2", "text": f"{card['title']}　{_fmt(card['score'])} 分"})
            if card["explanation"]:
                blocks.append({"type": "p", "text": card["explanation"]["text"]})
            # 毒性一格：只列有危险性说明的物质，每个物质最多三条，写成“物质：说明、说明；物质：……”
            hazards = "；".join(
                f"{h['name']}：{'、'.join(s['text'] for s in h['statements'][:3])}"
                for h in card["hazards"] if h["statements"]
            ) or "没有查到需要标注的危害"
            # 原料一格：这一步要准备的物料和它们的状态，写成“名字（可购买）”
            materials = "、".join(f"{m['name']}（{m['label']}）" for m in card["materials"]) or "—"
            blocks.append({
                "type": "table",
                "header": ["项目", "内容"],
                "rows": [
                    ["证据等级", f"{_evidence(card)}　·　置信度 {card['confidence_label']}"],
                    ["文献依据", card["evidence"]["basis"] or "—"],
                    ["反应条件", card["conditions"].get("text") or "没有条件记录"],
                    ["操作", _operation(card)],
                    ["产率", card["yield"]["label"]],
                    ["毒性", hazards],
                    ["原料", materials],
                    ["绿色化学", card["green"].get("basis") or "—"],
                ],
                "widths": [1, 6],
            })
            # 六个维度：权重写成百分数（0.3 → 30%），得分没有数据时显示“—”
            blocks.append({
                "type": "table",
                "header": ["维度", "权重", "得分", "依据"],
                "rows": [[d["label"], f"{d['weight'] * 100:g}%", _fmt(d["score"]), d["basis"]] for d in card["dimensions"]],
                "widths": [1.3, 0.7, 0.7, 5],
            })

    # ---- 三、文献列表：各步的先例合成一张表
    if "literature" in sections:
        blocks.append({"type": "h1", "text": "三、文献列表"})
        rows = []
        for card in cards:
            for precedent in card["precedents"]:
                # precedent["ref"] 形如 "P3"，[1:] 去掉开头的 P，只留数字，和说明文字里的 [3] 对得上
                rows.append([
                    f"第 {card['step_index']} 步", precedent["ref"][1:], precedent["doc_type_label"], precedent["document_id"],
                    _fmt(precedent["year"]), "同一反应" if precedent["exact"] else f"{precedent['similarity']:.2f}",
                    _fmt(precedent["yield_pct"], "%"), precedent["url"] or "—",
                ])
        if rows:
            blocks.append({"type": "table", "header": ["步骤", "编号", "类型", "文献号", "年份", "相似度", "产率", "链接"],
                           "rows": rows, "widths": [0.8, 0.5, 0.7, 1.6, 0.6, 0.8, 0.6, 3]})
        else:
            blocks.append({"type": "p", "text": "这条路线的各步都没有检索到文献先例。"})

    # ---- 四、试剂与物料清单
    if "materials" in sections:
        # 折算的目标产量（毫克）：用户指定了就用指定的，否则用这个规模的默认值（[scale]）
        amount = float(target_mg or scoring["scale"][scale])
        bom = bill_of_materials(result, route, amount)
        blocks.append({"type": "h1", "text": f"四、试剂与物料清单（按得到 {amount:g} mg 目标物折算）"})
        # 试剂没有用量：mmol 一格显示“—”，mg 一格写“参照文献当量”
        blocks.append({
            "type": "table",
            "header": ["物料", "角色", "用于", "分子量", "用量 / mmol", "用量 / mg", "来源"],
            "rows": [[r["name"], r["role"], "、".join(f"第 {s} 步" for s in r["steps"]), _fmt(r["mol_weight"]),
                      _fmt(r["mmol"]), _fmt(r["mg"]) if r["mg"] is not None else "参照文献当量", r["source"]]
                     for r in bom["rows"]],
            "widths": [2, 0.9, 1.2, 0.8, 1, 1.2, 0.9],
        })
        note = bom["note"]
        if bom["assumed_full_yield_steps"]:
            steps = "、".join(str(s) for s in bom["assumed_full_yield_steps"])
            note += f"第 {steps} 步没有产率记录，按 100% 计算，实际需要更多。"
        blocks.append({"type": "note", "text": note})

    # ---- 五、安全汇总
    if "safety" in sections:
        blocks.append({"type": "h1", "text": "五、安全汇总"})
        # seen：物质的 key14 -> {"hazard": 这个物质的危害信息, "steps": 用到它的步骤序号}。
        # 同一个物质在几步里出现只留一条。这里包括每一步的反应物、试剂和产物
        seen: dict[str, dict[str, Any]] = {}
        for card in cards:
            for hazard in card["hazards"]:
                entry = seen.setdefault(hazard["key14"], {"hazard": hazard, "steps": set()})
                entry["steps"].add(card["step_index"])
        rows = []
        for entry in seen.values():
            hazard = entry["hazard"]
            # 既没有危险性说明、也不在管制清单上、数据来源也不是“没有数据”的物质，不用列
            if not (hazard["statements"] or hazard["controlled"] or hazard["data_source"] == "none"):
                continue
            controlled = "、".join(f"{c['jurisdiction']} {c['category']}" for c in hazard["controlled"]) or "—"
            statements = "；".join(f"{s['code']} {s['text']}" for s in hazard["statements"][:5]) or "没有危害数据"
            rows.append([hazard["name"], "、".join(f"第 {s} 步" for s in sorted(entry["steps"])), statements, controlled])
        blocks.append({"type": "table", "header": ["物质", "用于", "GHS 危险性说明", "管制"], "rows": rows,
                       "widths": [1.5, 1, 4.2, 1.3]})
        # 特殊操作：各步命中的全部操作标签（M7），写成“第 N 步：标签（理由）”
        operations = [f"第 {c['step_index']} 步：{t['label']}（{t['reason']}）" for c in cards for t in c["conditions"].get("tags") or []]
        if operations:
            blocks.append({"type": "h2", "text": "特殊操作"})
            blocks.append({"type": "bullets", "items": operations})
        # data_source 是 "local" 表示危害分类取自内置试剂表（data/compounds.csv）；reviewed 是那一行有没有复核过
        unreviewed = any(e["hazard"]["data_source"] == "local" and not e["hazard"]["reviewed"] for e in seen.values())
        if unreviewed:
            blocks.append({"type": "note", "text": "部分危害分类来自内置试剂表的初稿，还没有经过复核，以 SDS 为准。"})

    # ---- 免责声明（不管选了哪些部分都有），以及这次分析的数据说明
    blocks.append({"type": "h1", "text": "免责声明"})
    blocks.append({"type": "p", "text": DISCLAIMER})
    for warning in result.get("warnings") or []:
        blocks.append({"type": "note", "text": f"数据说明：{warning}"})
    return blocks


# ---------------------------------------------------------------- PDF

# 已经注册好的 PDF 字体的名字。同一个进程里只找一次、注册一次，之后直接用（见 _pdf_font）。
_PDF_FONT: str | None = None
# 依次尝试的中文字体文件：Windows 的微软雅黑、黑体、宋体；macOS 的华文黑体、冬青黑体、Arial Unicode；
# Linux 的文泉驿微米黑、文泉驿正黑、Droid Sans Fallback。哪个文件存在并且嵌得进去，就用哪个。
_FONT_CANDIDATES = [
    r"C:\Windows\Fonts\msyh.ttc",
    r"C:\Windows\Fonts\simhei.ttf",
    r"C:\Windows\Fonts\simsun.ttc",
    "/System/Library/Fonts/STHeiti Light.ttc",
    "/System/Library/Fonts/Hiragino Sans GB.ttc",
    "/Library/Fonts/Arial Unicode.ttf",
    "/usr/share/fonts/truetype/wqy/wqy-microhei.ttc",
    "/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc",
    "/usr/share/fonts/truetype/droid/DroidSansFallbackFull.ttf",
]


def _pdf_font(font_path: str = "") -> str:
    """找一个能嵌进 PDF 的中文字体。服务器上没有时，用 REPORT_FONT_PATH 指定一个 TrueType 字体文件。

    参数：
        font_path  配置项 REPORT_FONT_PATH 的值；给了就最先试它。空字符串表示没配。
    返回：
        在 ReportLab 里注册好的字体名：嵌入成功时是 "SyntheriskCJK"，都不成功时是 "STSong-Light"。
    副作用：
        读字体文件，往 ReportLab 注册字体，并把结果记在模块级变量 _PDF_FONT 里。
        第一次调用之后就直接返回记下的名字，所以进程启动以后再改 font_path 不会生效。
    谁在用：
        本文件的 render_pdf()。
    """
    # global：下面要给模块级变量 _PDF_FONT 赋值
    global _PDF_FONT
    if _PDF_FONT:
        return _PDF_FONT
    # ReportLab 在函数里才导入：只有真的要出 PDF 时才加载它
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.pdfbase.ttfonts import TTFont

    # 先试配置里指定的字体，再按顺序试候选列表
    for path in ([font_path] if font_path else []) + _FONT_CANDIDATES:
        if os.path.exists(path):
            try:
                # .ttc 是把几个字体打包在一起的文件，subfontIndex=0 取其中第一个
                pdfmetrics.registerFont(TTFont("SyntheriskCJK", path, subfontIndex=0))
                _PDF_FONT = "SyntheriskCJK"
                return _PDF_FONT
            except Exception:  # noqa: BLE001 - 这个字体文件嵌不进去，换下一个
                continue
    # 兜底：不嵌字体，靠阅读器自带的中文字体显示，个别阅读器里可能显示不全
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    _PDF_FONT = "STSong-Light"
    return _PDF_FONT


def render_pdf(blocks: list[Block], font_path: str = "") -> bytes:
    """把一串块排成 PDF 文件（A4 纵向）。

    参数：
        blocks     build_blocks() 的返回值。
        font_path  配置项 REPORT_FONT_PATH（见 _pdf_font）。
    返回：
        PDF 文件的内容（bytes，以 %PDF 开头）。
    各种块怎么排：
        title / h1 / h2 / p / caption / note   各有一种段落样式（字号和颜色见下面的 styles）
        bullets     每项一段，前面加一个圆点
        image       按 width_mm 定宽度（不超过版心宽度），高度 = 宽度 × ratio；png 是 None 的跳过
        table       各列按 widths 的比例分版心宽度；单元格里的文字会自动换行；表格跨页时每页重复表头
        pagebreak   另起一页
    页面：A4，左右边距 18 毫米，上下边距 16 毫米；文档标题属性是“Syntherisk 合成路线报告”。
    副作用：
        第一次调用时读字体文件并注册字体（_pdf_font）。不写磁盘，PDF 生成在内存里。
    谁在用：
        api/reports.py 的 create_report()（F3，format 是 pdf 时）。
    """
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import Image, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    font = _pdf_font(font_path)
    # 基础段落样式：字号 9.5 磅，行距（leading）14.5 磅。
    # wordWrap="CJK" 是按中文的方式换行：中文的字与字之间没有空格，不这样设，长句子不会在中间断行
    base = ParagraphStyle("base", fontName=font, fontSize=9.5, leading=14.5, wordWrap="CJK")
    # 各种块的样式都从 base 派生，只改字号、行距、段前段后的间距（spaceBefore / spaceAfter，单位磅）和颜色
    styles = {
        "title": ParagraphStyle("title", parent=base, fontSize=20, leading=28, spaceAfter=6),
        "h1": ParagraphStyle("h1", parent=base, fontSize=14, leading=20, spaceBefore=12, spaceAfter=6, textColor=colors.HexColor("#0f5f66")),
        "h2": ParagraphStyle("h2", parent=base, fontSize=11.5, leading=17, spaceBefore=8, spaceAfter=4),
        "p": ParagraphStyle("p", parent=base, spaceAfter=5),
        "caption": ParagraphStyle("caption", parent=base, fontSize=9, textColor=colors.HexColor("#555555"), spaceBefore=6),
        "note": ParagraphStyle("note", parent=base, fontSize=8.5, leading=13, textColor=colors.HexColor("#666666"), spaceAfter=4),
        "cell": ParagraphStyle("cell", parent=base, fontSize=8.5, leading=12.5),
    }
    # 版心宽度 = A4 的宽度减去左右各 18 毫米的边距。ReportLab 的长度单位是磅，mm 是“1 毫米等于多少磅”
    page_width = A4[0] - 36 * mm
    # story：按先后顺序排进文档的内容（段落、表格、图片……），ReportLab 自己负责分页
    story: list[Any] = []
    for block in blocks:
        kind = block["type"]
        if kind in ("title", "h1", "h2", "p", "caption", "note"):
            # Paragraph 会把文字当成带标记的文本来解析（类似 HTML），所以先用 escape() 把 & < > 转义掉
            story.append(Paragraph(escape(block["text"]), styles[kind]))
        elif kind == "bullets":
            for item in block["items"]:
                story.append(Paragraph("• " + escape(item), styles["p"]))
        elif kind == "image" and block.get("png"):
            # 图片从内存里的字节读（io.BytesIO 把 bytes 包装成一个“文件”）；宽度不超过版心
            width = min(block["width_mm"] * mm, page_width)
            story.append(Image(io.BytesIO(block["png"]), width=width, height=width * block["ratio"]))
        elif kind == "table":
            # 各列的实际宽度 = 版心宽度 × 这一列的相对宽度 ÷ 相对宽度之和
            total = sum(block["widths"])
            widths = [page_width * w / total for w in block["widths"]]
            # 表头和各行合成一张二维表；每个单元格包成 Paragraph，长文字才会在格子里换行
            data = [[Paragraph(escape(str(cell)), styles["cell"]) for cell in row] for row in [block["header"], *block["rows"]]]
            # repeatRows=1：表格跨页时，每一页都重复第一行（表头）
            table = Table(data, colWidths=widths, repeatRows=1)
            # 表格样式。坐标是 (列, 行)，(0, 0) 是左上角，(-1, -1) 是右下角：
            # 全表画 0.4 磅的格线；第一行（表头）加底色；内容靠上对齐；上下各留 3 磅的内边距
            table.setStyle(TableStyle([
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#c9d3d6")),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e6f1f2")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("TOPPADDING", (0, 0), (-1, -1), 3),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ]))
            story.append(table)
            # 表格后面空 6 磅
            story.append(Spacer(1, 6))
        elif kind == "pagebreak":
            story.append(PageBreak())
    # PDF 写进内存里的缓冲区，不落盘；build() 做实际的排版
    buffer = io.BytesIO()
    SimpleDocTemplate(buffer, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=16 * mm,
                      bottomMargin=16 * mm, title="Syntherisk 合成路线报告").build(story)
    return buffer.getvalue()


# ---------------------------------------------------------------- Word


def render_docx(blocks: list[Block]) -> bytes:
    """把一串块排成 Word 文件（.docx）。

    参数：
        blocks  build_blocks() 的返回值。
    返回：
        .docx 文件的内容（bytes）。
    各种块怎么排：
        title       Word 的“标题”样式（add_heading 的 level=0）
        h1 / h2     “标题 1”“标题 2”样式
        p           普通段落
        caption / note   9 磅灰色小字
        bullets     每项一段，用“List Bullet”（项目符号列表）样式
        image       宽度取 width_mm，最多 160 毫米；高度按图片原来的比例；png 是 None 的跳过
        table       带格线的表格（“Table Grid”样式），后面空一段；块里的 widths 在 Word 里没有用到
        pagebreak   分页符
    字体：正文和各级标题都设成微软雅黑（Microsoft YaHei），正文 10.5 磅。
    文件里记的只是字体的名字，没有把字体本身放进去。
    副作用：没有。不写磁盘，文件生成在内存里。
    谁在用：
        api/reports.py 的 create_report()（F3，format 是 docx 时）。
    """
    # python-docx 在函数里才导入：只有真的要出 Word 时才加载它
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt, RGBColor

    document = Document()
    # 把用到的四种样式的字体都改成微软雅黑。Word 里西文和中文的字体是分开设的：
    # style.font.name 只管西文；中文（东亚文字）的字体要直接改样式的 XML，设 w:eastAsia 这个属性
    for name in ("Normal", "Title", "Heading 1", "Heading 2"):
        style = document.styles[name]
        style.font.name = "Microsoft YaHei"
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Microsoft YaHei")
    document.styles["Normal"].font.size = Pt(10.5)

    for block in blocks:
        kind = block["type"]
        if kind == "title":
            document.add_heading(block["text"], level=0)
        elif kind == "h1":
            document.add_heading(block["text"], level=1)
        elif kind == "h2":
            document.add_heading(block["text"], level=2)
        elif kind == "p":
            document.add_paragraph(block["text"])
        elif kind in ("caption", "note"):
            # 图题和小字说明：新起一段，里面放一段 9 磅、灰色（#666666）的文字
            run = document.add_paragraph().add_run(block["text"])
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
        elif kind == "bullets":
            for item in block["items"]:
                document.add_paragraph(item, style="List Bullet")
        elif kind == "image" and block.get("png"):
            # 只指定宽度，高度由 python-docx 按图片原来的比例算
            document.add_picture(io.BytesIO(block["png"]), width=Mm(min(block["width_mm"], 160)))
        elif kind == "table":
            # 先建只有一行（表头）的表格，再一行一行往下加
            table = document.add_table(rows=1, cols=len(block["header"]))
            table.style = "Table Grid"
            for cell, text in zip(table.rows[0].cells, block["header"]):
                cell.text = str(text)
            for row in block["rows"]:
                for cell, text in zip(table.add_row().cells, row):
                    cell.text = str(text)
            # 表格后面加一个空段落，和下面的内容隔开
            document.add_paragraph()
        elif kind == "pagebreak":
            document.add_page_break()
    # 存进内存里的缓冲区，取出字节返回
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


# ---------------------------------------------------------------- SDF + RXN


def export_structures_zip(result: dict[str, Any], route: dict[str, Any]) -> bytes:
    """一个 zip：route.sdf 里是路线上的全部分子，step_N.rxn 是每一步的反应文件。

    这是 F2 里“sdf”格式实际给出的文件。SDF 和 RXN 是化学软件之间交换分子结构、反应的常用文件格式。

    参数：
        result  分析结果（用 molecules 和 reactions 两张表）。
        route   要导出的路线（用它的 tree、leaf_keys、step_keys、name）。
    返回：
        zip 文件的内容（bytes）。里面有：
        route.sdf     SDF 格式（一个文件里放多个分子，每个分子带二维坐标和几个属性）。
                      包含路线树上的全部分子和各步的试剂，每个分子带这些属性：
                          标题行（_Name）  名字；没有名字时是分子式
                          Role            target（目标分子）/ intermediate（中间体）/
                                          starting_material（起始原料）/ reagent（试剂）
                          SMILES、InChIKey
                          Availability    可得性的中文标签（库存 / 可购买 / 买不到）；没有时是空的
        step_N.rxn    MDL RXN 格式，每一步一个：反应物 → 产物，不含试剂。N 是步骤序号，从 1 开始。
                      某一步生成不了反应文件时，就没有那一步的文件
        README.txt    路线名、每个文件是什么，以及免责声明
    副作用：没有。不写磁盘，zip 生成在内存里。
    谁在用：
        api/reports.py 的 create_report()（F3，format 是 sdf 时；存下来的文件扩展名是 zip）。
    """
    molecules, reactions = result["molecules"], result["reactions"]
    leaves = set(route["leaf_keys"])
    target_key = route["tree"]["key14"]

    # roles：分子的 key14 -> 它在这条路线里的角色。
    # iter_mol_nodes() 遍历路线树上的每个分子节点：根是目标分子，叶子是起始原料，其余是中间体
    roles: dict[str, str] = {}
    for node in iter_mol_nodes(route["tree"]):
        key = node["key14"]
        roles[key] = "target" if key == target_key else ("starting_material" if key in leaves else "intermediate")
    # 各步的试剂。setdefault：已经有角色的分子（比如它同时是某一步的原料）保持原来的角色
    for step_key in route["step_keys"]:
        for key in reactions[step_key]["reagents"]:
            roles.setdefault(key, "reagent")

    # SDWriter 是 RDKit 写 SDF 的工具，这里让它写进内存里的字符串缓冲区
    sdf = io.StringIO()
    writer = Chem.SDWriter(sdf)
    for key, role in roles.items():
        entry = molecules[key]
        # to_mol() 返回的分子对象是带缓存、大家共用的；Chem.Mol(...) 复制一份再改，不动缓存里的那一份
        mol = Chem.Mol(chem.to_mol(entry["smiles"]))
        # SMILES 没有坐标，算一套二维坐标，别的软件打开时才画得出来
        AllChem.Compute2DCoords(mol)
        # SetProp 设的属性会写进 SDF；"_Name" 比较特殊，它是分子记录的第一行（标题行）
        mol.SetProp("_Name", entry.get("name") or entry["formula"])
        mol.SetProp("Role", role)
        mol.SetProp("SMILES", entry["smiles"])
        mol.SetProp("InChIKey", entry["inchikey"])
        mol.SetProp("Availability", (entry.get("availability") or {}).get("label", ""))
        writer.write(mol)
    # close() 把还没写出去的内容写完，之后 sdf.getvalue() 才是完整的
    writer.close()

    buffer = io.BytesIO()
    # ZIP_DEFLATED：压缩存放。writestr(文件名, 内容) 往压缩包里加一个文件
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("route.sdf", sdf.getvalue())
        for index, step_key in enumerate(route["step_keys"], 1):
            reaction = reactions[step_key]
            # reaction_rxn_block()（m01_structure）：由反应物和产物的 SMILES 生成 RXN 文件的文本
            block = chem.reaction_rxn_block(
                [molecules[k]["smiles"] for k in reaction["reactants"]], molecules[reaction["product"]]["smiles"]
            )
            if block:
                archive.writestr(f"step_{index}.rxn", block)
        # README：每一步的文件名和反应名各一行，最后是免责声明
        steps = "\n".join(f"step_{i}.rxn  第 {i} 步 {reactions[k]['name']}" for i, k in enumerate(route["step_keys"], 1))
        archive.writestr("README.txt", f"路线：{route['name']}\nroute.sdf  路线上的全部分子（Role 字段是它在路线里的角色）\n{steps}\n\n{DISCLAIMER}\n")
    return buffer.getvalue()


# ---------------------------------------------------------------- D4 对比表


def _comparison_rows(comparison: dict[str, Any]) -> tuple[list[str], list[list[Any]]]:
    """把对比结果排成一张二维表：一行一个指标，一列一条路线。

    参数：
        comparison  m10_scoring.compare_routes() 的返回值（{"dimension_labels", "routes"}）。
    返回：
        (header, rows)
        header  表头：["指标", 路线 1 的名字, 路线 2 的名字, ...]
        rows    各行，每行的第一格是指标名，后面是各条路线的值（可能是 None）。依次是：
                综合分、步数、文献覆盖 %、库存覆盖 %、特殊操作、原料成本、置信度；
                六个维度各一行（“<维度名>（各步平均）”）；
                再按步骤序号一行一步（“第 N 步”）：格子里是那条路线这一步的反应名，
                各条路线共有的步骤后面加“（共同步骤）”；路线比较短、没有这一步时格子是空的。
    谁在用：
        本文件的 comparison_xlsx() 和 comparison_csv()。
    """
    routes = comparison["routes"]
    # [..., *列表] 是把列表里的元素摊开放进来：第一格是指标名，后面每条路线一格
    header = ["指标", *[route["name"] for route in routes]]
    labels = comparison["dimension_labels"]
    rows: list[list[Any]] = [
        ["综合分", *[r["score"] for r in routes]],
        ["步数", *[r["n_steps"] for r in routes]],
        ["文献覆盖 %", *[r["literature_coverage"] for r in routes]],
        ["库存覆盖 %", *[r["inventory_coverage"] for r in routes]],
        ["特殊操作", *[r["special_operations"] for r in routes]],
        ["原料成本", *[r["material_cost"] for r in routes]],
        ["置信度", *[r["confidence"] for r in routes]],
    ]
    # 六个维度各一行：radar 是这条路线各步在每个维度上的平均分
    for key, label in labels.items():
        rows.append([f"{label}（各步平均）", *[r["radar"].get(key) for r in routes]])
    # 步骤：行数取最长那条路线的步数；短的路线后面的格子留空
    longest = max((len(r["steps"]) for r in routes), default=0)
    for position in range(longest):
        cells = []
        for route in routes:
            step = route["steps"][position] if position < len(route["steps"]) else None
            cells.append("" if step is None else f"{step['name']}{'（共同步骤）' if step['shared'] else ''}")
        rows.append([f"第 {position + 1} 步", *cells])
    return header, rows


def comparison_xlsx(comparison: dict[str, Any], weights: dict[str, float]) -> bytes:
    """把路线对比导出成 Excel 文件（.xlsx）。

    参数：
        comparison  m10_scoring.compare_routes() 的返回值。
        weights     这次对比用的六个权重（键是维度的英文键）。
    返回：
        .xlsx 文件的内容（bytes）。里面有两张表：
        “路线对比”  _comparison_rows() 的表：第一行是表头（加粗），第一列是指标名（列宽 26）；没有值的格子是空的
        “权重”      两列：维度的中文名、权重
    副作用：没有。不写磁盘，文件生成在内存里。
    谁在用：
        api/reports.py 的 export_comparison()（D4，format 是 xlsx 时）。
    """
    # openpyxl 在函数里才导入：只有真的要出 Excel 时才加载它
    from openpyxl import Workbook
    from openpyxl.styles import Font

    header, rows = _comparison_rows(comparison)
    workbook = Workbook()
    # 新建的工作簿自带一张表，直接拿来用
    sheet = workbook.active
    sheet.title = "路线对比"
    sheet.append(header)
    # sheet[1] 是第一行的全部单元格：表头加粗
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in rows:
        sheet.append(["" if value is None else value for value in row])
    sheet.column_dimensions["A"].width = 26
    # 第二张表：这次对比用的权重
    weight_sheet = workbook.create_sheet("权重")
    weight_sheet.append(["维度", "权重"])
    for key, label in comparison["dimension_labels"].items():
        weight_sheet.append([label, float(weights.get(key, 0))])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def comparison_csv(comparison: dict[str, Any]) -> bytes:
    """把路线对比导出成 CSV 文件。

    参数：
        comparison  m10_scoring.compare_routes() 的返回值。
    返回：
        CSV 文件的内容（bytes）：_comparison_rows() 的表，逗号分隔，没有值的格子是空的。
        编码是带 BOM 的 UTF-8。CSV 里没有权重那张表。
    副作用：没有。
    谁在用：
        api/reports.py 的 export_comparison()（D4，format 是 csv 时）。
    """
    header, rows = _comparison_rows(comparison)
    # csv.writer 负责给含有逗号、引号的格子加引号；先写进内存里的字符串缓冲区
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(header)
    writer.writerows([["" if value is None else value for value in row] for row in rows])
    return buffer.getvalue().encode("utf-8-sig")  # 带 BOM，Excel 直接打开不乱码
